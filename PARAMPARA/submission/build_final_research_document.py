from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.enum.dml import MSO_THEME_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "final"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "PARAMPARA_SIH_Submission_Research_and_Evidence_Dossier.docx"

NAVY = "17324D"
BLUE = "DCEAF5"
PALE = "F4F7FA"
GOLD = "B88A3B"
GRAY = "596773"
WHITE = "FFFFFF"
BLACK = "000000"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.6)
styles["Normal"].font.color.rgb = RGBColor.from_string("20262C")
styles["Normal"].paragraph_format.space_after = Pt(5.5)
styles["Normal"].paragraph_format.line_spacing = 1.08

for name, size, before, after in [
    ("Title", 28, 0, 10), ("Subtitle", 12, 0, 10),
    ("Heading 1", 19, 12, 7), ("Heading 2", 13, 10, 4),
    ("Heading 3", 10.5, 8, 3),
]:
    st = styles[name]
    st.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
styles["Heading 1"].paragraph_format.page_break_before = True

if "Caption" in styles:
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string(GRAY)
    styles["Caption"].paragraph_format.space_after = Pt(7)

for style_name in ["TOC 1", "TOC 2", "TOC 3"]:
    if style_name not in styles:
        styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def keep_table_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)

def add_table(headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    keep_table_row_together(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        set_cell_shading(c, NAVY)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        r.bold = True
        r.font.color.rgb = RGBColor(255,255,255)
        r.font.size = Pt(font_size)
        if widths: c.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        added_row = table.add_row()
        keep_table_row_together(added_row)
        cells = added_row.cells
        for i, value in enumerate(row):
            c = cells[i]
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1: set_cell_shading(c, PALE)
            if widths: c.width = Inches(widths[i])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
    return table

def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def add_numbered(items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{number}. ").bold = True
        p.add_run(item)

def add_caption(text):
    p = doc.add_paragraph(text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def add_image(path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(caption)

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve"); instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])

def add_toc():
    p = doc.add_paragraph()
    r = p.add_run()
    fldChar = OxmlElement("w:fldChar"); fldChar.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:t"); fldChar3.text = "Right-click and update field to refresh the table of contents."
    fldChar4 = OxmlElement("w:fldChar"); fldChar4.set(qn("w:fldCharType"), "end")
    r._r.extend([fldChar, instrText, fldChar2, fldChar3, fldChar4])

def new_section(title, intro=None):
    doc.add_heading(title, level=1)
    if intro:
        p = doc.add_paragraph(intro)
        p.paragraph_format.space_after = Pt(8)

def para(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

# Header and footer
header = sec.header.paragraphs[0]
header.text = "PARAMPARA  |  SIH RESEARCH AND EVIDENCE DOSSIER"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(7.5); run.font.color.rgb = RGBColor.from_string(GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("PARAMPARA  •  Evidence controlled edition  •  ")
r.font.size = Pt(7.5); r.font.color.rgb = RGBColor.from_string(GRAY)
add_page_number(footer)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(52)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PARAMPARA")
r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(37); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Research Engineering and Prototype Evidence Dossier")
p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Smart India Hackathon 2026  |  Hardware  |  Heritage and Culture")

cover_img = ROOT / "cad" / "renders" / "sleeve_iso.png"
add_image(cover_img, 5.9, "Digital CAD concept of the full sleeve. This is not a photograph of an assembled physical prototype.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
r = p.add_run("Problem Statement 26214")
r.bold = True; r.font.size = Pt(12)
p = doc.add_paragraph("Student Innovation  Ideas that showcase the rich cultural heritage and traditions of India")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_table(["Submission field", "Entry"], [
    ("Team name", "TO BE FILLED BEFORE SUBMISSION"),
    ("Team ID", "TO BE FILLED BEFORE SUBMISSION"),
    ("Institute", "TO BE FILLED BEFORE SUBMISSION"),
    ("Team leader and members", "TO BE FILLED BEFORE SUBMISSION"),
    ("Mentor", "TO BE FILLED BEFORE SUBMISSION"),
    ("Version", "1.0  •  5 October 2026"),
], widths=[1.75, 4.7], font_size=8.4)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Repository  ").bold = True
add_hyperlink(p, "github.com/HarshDubey23/winner", "https://github.com/HarshDubey23/winner")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Interactive circuit  ").bold = True
add_hyperlink(p, "Wokwi project 477036096886805505", "https://wokwi.com/projects/477036096886805505")

doc.add_page_break()
doc.add_heading("Document purpose and submission status", level=1)
para("This dossier presents PARAMPARA as a technically testable proposal for preserving embodied craft knowledge. It combines an open palm sensing and haptic sleeve, synchronized tool events, a consent governed Skill Fingerprint, adaptive cue fading, and device off assessment. Tabla is the first validation case. Kathputli puppetry and handloom are planned extensions that require their own field validation.")
para("The main conclusion is deliberately bounded. The repository demonstrates CAD, software logic, virtual electronics, a browser practice prototype, reproducible analysis on public electronic drum MIDI, and a professional concept film. It does not yet demonstrate a manufactured sleeve, measured wearable accuracy, real tabla master identification, comfort, battery life, or improved human learning. Those claims are assigned to a staged validation plan.")
add_table(["Evidence state", "Meaning in this dossier"], [
    ("Implemented", "Code or digital artifact exists and can be inspected or run."),
    ("Tested in software", "Automated or simulated checks passed in the stated environment."),
    ("Designed", "CAD, architecture or protocol exists; physical performance is unverified."),
    ("Proposed", "Future implementation or study; no result is claimed."),
], widths=[1.5, 5.0])
doc.add_heading("SIH compliance note", level=2)
para("The latest official college guideline located during preparation is the SIH 2024 College SPOC guideline. It lists novelty, complexity, clarity in the prescribed format, feasibility, practicability, sustainability, scale of impact, user experience and future work progression as evaluation criteria [1]. This dossier maps the project to those criteria. The current SIH portal, official presentation template, file size limits and institute SPOC instructions must be checked before upload; they override this supporting dossier.")
doc.add_heading("Contents", level=2)
add_toc()

new_section("1 Executive summary")
para("India's craft traditions carry knowledge in timing, touch, force, coordination and movement. Video can record what a movement looks and sounds like, but it does not directly encode every joint trajectory, finger event or cue timing that a learner may need. PARAMPARA investigates whether a consenting master's movement and tool events can be captured as a reference and converted into progressively reduced guidance that supports independent practice.")
doc.add_heading("Proposed contribution", level=2)
add_bullets([
    "An open palm wearable architecture that keeps fingertips and the palm free while proposing nine motion sites and eight haptic cue sites per arm.",
    "A synchronized body plus tool representation, termed a Skill Fingerprint, evaluated across sessions and instruments to reduce identity and equipment confounds.",
    "An adaptive teaching loop that schedules anticipatory cues, reduces guidance in bounded steps, inserts regular no cue checks and evaluates delayed device off performance.",
    "A consent and provenance layer intended to preserve master attribution, permitted use, export and deletion choices.",
    "A staged evidence ladder from digital prototypes to one channel bench validation, wearable validation and controlled human learning studies."
])
doc.add_heading("Current proof package", level=2)
add_table(["Artifact", "Current evidence", "Boundary"], [
    ("3D equipment", "8 GLB scenes, 9 viewer views, 20 STEP and STL housing files", "Packaging intent; no physical fit approval"),
    ("Firmware core", "27 host checks plus 1,425 generated scoring, timing and fade checks", "Mock host execution; not full target hardware"),
    ("Virtual circuit", "Compiled and executed ESP32 Wokwi sketch with retained logs", "MPU6050 and LEDs substitute proposed devices"),
    ("Phone prototype", "On time, 60 ms late and silent input scenarios with CSV export", "Synthetic profiles and browser timing"),
    ("Data experiment", "Held session and same groove experiments on public electronic drum MIDI", "Not tabla, wearables or skill quality"),
    ("Communication", "4 minute 10 second jury film and 72 second technical video", "Concept and digital prototype evidence"),
], widths=[1.25, 3.1, 2.2], font_size=7.9)

new_section("2 Problem definition and need")
doc.add_heading("2.1 Problem hypothesis", level=2)
para("The project starts with a researchable gap: conventional audio and video may not preserve all of the movement and timing information that supports craft instruction. PARAMPARA asks whether synchronized body and tool recordings contain repeatable, teachable information beyond the visible and audible record, and whether fading cues based on cue free checks improves later unaided performance.")
doc.add_heading("2.2 Primary users and stakeholders", level=2)
add_table(["Stakeholder", "Need", "Design implication"], [
    ("Master practitioners", "Control, attribution and faithful representation", "Consent, provenance, deletion and master review before release"),
    ("Learners", "Clear feedback without permanent dependence", "Short cues, bounded fading, no cue checks and delayed retention tests"),
    ("Institutions and gurukuls", "Reusable teaching support and evidence", "Offline operation, cohort analytics and exportable reports"),
    ("Craft organizations", "Preservation without extraction", "Community governance and benefit sharing"),
    ("Researchers", "Reproducible protocols and honest metrics", "Predefined tests, raw logs and claim boundaries"),
], widths=[1.35, 2.5, 2.85], font_size=8.0)
doc.add_heading("2.3 Focused first use case", level=2)
para("Tabla is the flagship validation case because timing and bilateral coordination can be linked to discrete tool events. The first study should use one defined Teentaal exercise, a small number of movement sites and two drum event channels. Expansion to a full sleeve should follow only after repeatable one channel timing is measured.")
doc.add_heading("2.4 Research questions", level=2)
add_numbered([
    "Identity: can the representation distinguish consenting performers across recording days and swapped instruments?",
    "Teaching: does adaptive fading improve delayed, unaided performance compared with fixed cues, always on cues and audio video practice?",
    "Practicality: can the system operate without materially changing technique, comfort or attention?",
    "Governance: can a master inspect, approve, restrict and revoke use of their recordings and derived model?"
])

new_section("3 Research foundation and prior art")
para("PARAMPARA combines established components in a specific validation and governance design. Wearable IMUs, vibrotactile guidance and motor learning feedback have prior art. The defensible contribution is the integration of master consent, synchronized body and tool events, cross session and cross instrument validation, adaptive fading and device off assessment for craft learning.")
add_table(["Area", "What prior work establishes", "Gap addressed by PARAMPARA"], [
    ("Haptic skill learning", "Tactile feedback can support force and music skill tasks in controlled settings [6][7].", "Benefit must be tested for this craft, cue layout and curriculum."),
    ("Wearable motion capture", "Small inertial sensors can capture orientation and movement features.", "Repeatability after refitting and across tools remains essential."),
    ("Music datasets", "The Groove MIDI Dataset provides human timing and velocity for electronic drum performances [8].", "It is a method test, not evidence about tabla or wearables."),
    ("Adaptive guidance", "Guidance can be reduced as performance improves.", "Regular no cue checks and delayed retention reduce false confidence."),
    ("Traditional knowledge governance", "CARE principles emphasize collective benefit, authority, responsibility and ethics [12].", "Consent and permitted use should accompany every Skill Fingerprint."),
], widths=[1.35, 2.65, 2.65], font_size=7.9)
doc.add_heading("3.1 Novelty statement", level=2)
para("PARAMPARA does not claim that haptics, IMUs, gesture recognition or digital archives are individually new. Its novelty claim is the measurable combination: master governed capture, body and tool synchronization, confound controlled fingerprint tests, adaptive cue fading with bounded updates, and assessment that explicitly removes guidance.")
doc.add_heading("3.2 Intellectual property position", level=2)
para("No patentability or freedom to operate opinion has been completed. Before commercialization, the team should conduct a structured patent search covering wearable haptics, gesture teaching, skill transfer and motion signature systems. Code and documentation licences must be declared. Traditional knowledge and performer rights require governance beyond conventional software ownership.")

new_section("4 Solution overview")
add_image(ROOT / "figures" / "v4_pipeline.png", 6.55, "Figure 1  Proposed end to end workflow from master capture to learner assessment.")
doc.add_heading("4.1 Core workflow", level=2)
add_numbered([
    "Obtain informed consent and record permitted uses before capture.",
    "Record synchronized movement and tool events with quality checks and calibration metadata.",
    "Segment a phrase and extract timing, movement and relative force features.",
    "Validate repeatability on held out sessions and, where possible, swapped instruments.",
    "Publish only an approved reference package with provenance and permitted use fields.",
    "Guide the learner through audio, visual and anticipatory tactile cues.",
    "Reduce cue density only at predefined check cycles and evaluate without cues.",
    "Measure delayed, device off performance before claiming learning benefit."
])
doc.add_heading("4.2 Three craft modules", level=2)
add_table(["Module", "Primary events", "Proposed sensing", "Status"], [
    ("Tabla", "Bol timing, hand assignment and relative stroke strength", "Sleeves plus drum pickups", "Flagship validation case"),
    ("Kathputli", "Finger pulls, wrist paths and puppet response", "Sleeves plus puppet IMU or line sensing", "Planned extension"),
    ("Handloom", "Hand foot coordination, shuttle, beater and cycle timing", "Sleeves plus switches and beater sensing", "Planned extension"),
], widths=[1.1, 2.0, 2.15, 1.35], font_size=7.8)
add_image(ROOT / "figures" / "v5_puppet_loom_kits.png", 6.55, "Figure 2  Digital concept models for Kathputli and handloom tool sensing. These extensions have not been physically validated.")

new_section("5 System architecture")
add_image(ROOT / "figures" / "v5_block_diagram.png", 6.55, "Figure 3  Proposed electronic architecture and interface map.")
add_table(["Layer", "Function", "Proposed implementation"], [
    ("Capture", "Measure body movement and tool events", "Nine BMI270 sites per arm plus craft specific event sensors"),
    ("Routing", "Resolve repeated I2C addresses", "TCA9548A switched bus branches"),
    ("Hub", "Timestamp, buffer, store and communicate", "ESP32 S3 MINI 1, storage and power management"),
    ("Haptics", "Deliver local anticipatory cues", "Eight LRA sites driven through DRV2605L"),
    ("Analysis", "Segment and compare practice phrases", "Timing and movement features with held out evaluation"),
    ("Teaching", "Schedule and fade cues", "Bounded guidance updates and regular no cue cycles"),
    ("Governance", "Control use and provenance", "Consent record, attribution, export and deletion rules"),
], widths=[1.0, 2.25, 3.35], font_size=8.0)
doc.add_heading("5.1 Proposed sleeve site map", level=2)
para("One sleeve proposes nine motion sites at five fingers, the back of hand, forearm or wrist, upper arm and shoulder. Eight haptic sites are proposed at five fingers, wrist, elbow and shoulder. Fingertips and the palm remain free. A second sleeve supports bilateral practice.")
add_image(ROOT / "figures" / "v5_sleeve_views.png", 6.4, "Figure 4  Multi view CAD of the proposed open palm sleeve.")
doc.add_heading("5.2 Design inconsistency to resolve", level=2)
para("The CAD hub includes an additional BMI270 block while the firmware site map defines nine sites without a separate hub sensor. Before schematic freeze, the team must decide whether the hub sensor replaces a site, becomes a tenth site or is removed. CAD, firmware, wiring and BOM must then be updated together.")

new_section("6 Mechanical and electronic design")
add_image(ROOT / "figures" / "v5_pods_exploded.png", 6.45, "Figure 5  Exploded digital housings for the finger, haptic, hand board and hub modules.")
add_table(["Housing", "Nominal digital envelope", "Validation still required"], [
    ("Finger IMU ring", "13 × 11 × 4.4 mm", "Fit range, rotation, edge comfort and cable relief"),
    ("Haptic pod", "12 mm diameter × 5.2 mm", "Skin coupling, onset, audibility and attachment"),
    ("Back of hand board", "42 × 32 × 8 mm", "Clearance, flex, sweat protection and mass"),
    ("Joint pod", "28 × 20 × 8 mm", "Range of motion and migration"),
    ("Forearm hub", "62 × 44 × 16.4 mm", "Battery safety, thermal behavior and balance"),
], widths=[1.55, 1.7, 3.2], font_size=8.1)
doc.add_heading("6.1 Provisional one arm purchasing scope", level=2)
para("The provisional architecture requires nine motion sensors, eight haptic drivers, eight LRAs, one ESP32 S3 hub, I2C switches, local storage, protected battery and charging circuitry, cables, connectors, straps and printed housings. Tool sensors and a second sleeve are additional. Counts and prices must be updated after schematic reconciliation and dated supplier quotations.")
doc.add_heading("6.2 Engineering constraints", level=2)
add_bullets([
    "Sampling rate, timestamp alignment, sensor filter delay and actuator onset are different quantities and must be measured separately.",
    "A 200 Hz sensor produces a 5 ms sample interval; interpolation does not create faster physical samples.",
    "DRV2605L supports LRA and ERM actuation, but an LED in simulation does not demonstrate vibration amplitude, resonant tuning or rise time [4][5].",
    "The BMI270 requires a correct initialization sequence and configuration data; CAD placement alone does not demonstrate a working interface [3].",
    "The power budget must cover simultaneous sensing, storage, radio bursts and multiple motor cues, including brownout margins."
])

new_section("7 Adaptive teaching and Skill Fingerprint")
doc.add_heading("7.1 Skill Fingerprint", level=2)
para("A Skill Fingerprint is the proposed representation of timing, movement and relative force patterns associated with a consenting master performing a defined phrase. It is a research hypothesis, not a certified measure of artistic quality. A valid fingerprint must generalize to a new recording day and, where relevant, another instrument. Data splitting must prevent the same performance or near duplicate phrase from appearing in training and test sets.")
add_image(ROOT / "figures" / "confound_test.png", 5.75, "Figure 6  Confound controlled validation design using held sessions and instruments.")
doc.add_heading("7.2 Fade controller", level=2)
para("The implemented controller schedules a check when the cycle index modulo four equals zero. At each check, target guidance is clamp one minus score divided by 0.85 between zero and one. Guidance changes by no more than 0.20 at a check. Cue density moves through every beat, every second beat, every fourth beat, first beat only and none.")
add_image(ROOT / "submission" / "docs" / "figures" / "controller-trace.png", 5.75, "Figure 7  Scripted controller trace used to exercise guidance states. It is not a predicted learning curve.")
doc.add_heading("7.3 Scoring limitations", level=2)
para("The current core combines hand correctness and timing. It does not fully score tabla quality and does not explicitly penalize every extra tap. Before a human study, the team must fix definitions for extra strokes, wrong bols, amplitude, phrase boundaries, duplicate taps and permissible expressive timing. The scoring specification must be frozen before evaluation.")

new_section("8 Digital prototype and evidence")
doc.add_heading("8.1 Interactive 3D equipment", level=2)
para("The browser viewer presents the sleeve, hand close up, exploded modules and the three craft scenes. It supports rotation, zoom, labels, part selection, auto rotation and image export. Geometry is derived from repository CAD. It demonstrates design intent and communication quality, not manufacturability or comfort.")
add_image(ROOT / "showcase" / "evidence" / "hand-presentation.png", 6.25, "Figure 8  Interactive equipment viewer rendered from repository geometry.")
doc.add_heading("8.2 Virtual ESP32 electronics", level=2)
para("The public Wokwi project compiles and runs an ESP32 sketch. A virtual MPU6050 provides an acceleration input, blue and orange LEDs represent motion detection and cue output, and a button represents manual taps. The simulator demonstrates a one channel control flow and safety inhibition on sensor read failure. It does not demonstrate the proposed BMI270 array or real LRAs.")
add_image(ROOT / "submission" / "results" / "wokwi-running.png", 6.05, "Figure 9  Running Wokwi circuit used for virtual electronics evidence.")
add_table(["Condition", "Observed software output", "Interpretation"], [
    ("Default synthetic input 60 ms late", "Score 0.7500; guidance 0.80 at cycle 0 and 0.60 at cycle 4", "Controller behavior under a scripted offset"),
    ("On time synthetic input", "Score 1.0000; guidance eventually reaches 0.00", "Successful fixture can remove cues"),
    ("Synthetic taps disabled", "Score 0.0000; guidance remains 1.00", "Silence does not appear successful"),
], widths=[1.85, 2.45, 2.25], font_size=7.9)
doc.add_heading("8.3 Phone practice prototype", level=2)
para("PARAMPARA Lite implements a browser based rhythm practice loop with synthetic master profiles, cue fading, cue free checks and CSV export. On supported Android browsers, vibration can provide a simple haptic demonstration. Its timing depends on browser and device behavior and must not be presented as wearable latency.")
add_image(ROOT / "figures" / "v6_lite_screenshot.png", 5.7, "Figure 10  PARAMPARA Lite phone practice interface.")

new_section("9 Verification results and accuracy language")
doc.add_heading("9.1 Reproducible software checks", level=2)
add_table(["Test group", "Result", "What it supports", "What it does not support"], [
    ("Firmware host suite", "27 checks passed", "Mock routing, device behavior, fade rules, scoring and clock synchronization", "Physical ESP32 sleeve accuracy"),
    ("Generated benchmark", "1,000 scoring plus 425 timing and fade checks; zero failures", "Implemented controller matches expected fixture outputs", "Sensor accuracy or learning effectiveness"),
    ("Phone scenarios", "On time approximately 100 percent; 60 ms late approximately 75 percent; silent 0", "Scripted input discrimination", "Human performance percentage"),
    ("Wokwi scenarios", "On time, late and silent behaviors logged", "Virtual circuit control path", "Real actuator onset or electrical reliability"),
], widths=[1.35, 1.35, 2.25, 1.75], font_size=7.4)
doc.add_heading("9.2 Public data method experiment", level=2)
para("The Groove MIDI Dataset experiment asks whether timing and velocity features contain performer information across held sessions. Repository results report 47.6 percent balanced accuracy across seven electronic drum performers. In a controlled same groove subset with four performers, timing and velocity features achieved 63.125 percent balanced unit accuracy against a 25 percent chance level, while note identity alone achieved about 26 percent. These results motivate the method. They do not validate tabla, the wearable, a master fingerprint or artistic quality.")
add_image(ROOT / "submission" / "docs" / "figures" / "gmd-ablation.png", 4.6, "Figure 11  Repository experiment comparing feature groups on a same groove electronic drum subset.")
doc.add_heading("9.3 Correct use of the word accuracy", level=2)
add_table(["Acceptable statement", "Avoid"], [
    ("The generated software fixture passed 1,425 checks.", "The device is 100 percent accurate."),
    ("Held session balanced accuracy was 47.6 percent on electronic drum MIDI.", "PARAMPARA identifies tabla masters with 47.6 percent accuracy."),
    ("The Wokwi sketch responded correctly in retained scripted scenarios.", "The wearable is electronically validated."),
    ("The pilot protocol will test delayed unaided performance.", "The system improves learning or retention."),
], widths=[3.25, 3.25], font_size=8.0)

new_section("10 Validation plan")
para("The next work should reduce risk in sequence. A photographed one sensor to controller to motor bench demonstration is more valuable than expanding unverified features. Every experiment should define pass rules before data collection and retain failed as well as successful runs.")
add_table(["Stage", "Question", "Method", "Exit criterion"], [
    ("H1 one channel", "Can one sensor and one LRA operate with stable timing?", "Timestamped bench rig with oscilloscope or logic analyzer", "Measured latency distribution, no resets and safe temperature"),
    ("H2 routing", "Can repeated addresses be polled reliably?", "Scale switched bus while logging rate and errors", "Required channel rate with bounded error rate"),
    ("H3 actuation", "Are cues perceivable and localized?", "Bench acceleration plus blinded perception test", "Predefined perception and confusion thresholds"),
    ("E0 invasiveness", "Does the sleeve alter technique?", "Within participant comparison with counterbalanced order", "Equivalence margin fixed in advance"),
    ("E1 fingerprint", "Does identity survive day and instrument changes?", "Grouped cross validation by participant, session and instrument", "Predefined balanced accuracy and confidence interval"),
    ("E2 cue timing", "Are cues perceived before the movement window?", "Timestamped perception and onset study", "Predefined onset and perception window"),
    ("E3 cue reading", "Can users map cues to actions?", "Blinded cue discrimination task", "Predefined class accuracy"),
    ("E4 device off", "Does guidance reduce dependence?", "Regular cue free probe cycles", "Improvement without active guidance"),
    ("E5 learning pilot", "Does fading improve retention?", "Randomized groups and delayed test", "Pre-registered primary outcome and effect estimate"),
], widths=[1.05, 1.55, 2.55, 1.45], font_size=7.25)
doc.add_heading("10.1 Proposed pilot comparison", level=2)
para("A feasible early study can compare adaptive fading, always on tactile cues and audio video practice. Use a baseline test, repeated training blocks, immediate device off test and delayed test after a fixed interval. Randomize group allocation where possible, blind the scorer, report uncertainty and avoid selecting only favorable participants or phrases.")
doc.add_heading("10.2 Primary metrics", level=2)
add_bullets([
    "Timing error distribution in milliseconds rather than a single best case value.",
    "Correct hand and bol sequence, false positives, missed strokes and phrase completion.",
    "Cue onset latency, jitter, localization accuracy and perceived comfort.",
    "Unaided immediate and delayed performance relative to baseline and control groups.",
    "Dropout, adverse events, sleeve migration and calibration failure rates."
])

new_section("11 Feasibility implementation and cost")
doc.add_heading("11.1 Build sequence", level=2)
add_table(["Phase", "Duration target", "Deliverable"], [
    ("A architecture freeze", "Week 1", "Reconciled site map, schematic, BOM and test plan"),
    ("B one channel bench", "Weeks 2 to 3", "Sensor, controller and LRA with measured timing"),
    ("C multi channel rig", "Weeks 4 to 6", "Switched buses, synchronized tool event and power log"),
    ("D wearable alpha", "Weeks 7 to 10", "Printed housings, straps, safety checks and data capture"),
    ("E master capture pilot", "Weeks 11 to 14", "Consented recordings and cross session analysis"),
    ("F learner pilot", "Weeks 15 to 20", "Device off and delayed retention results"),
], widths=[1.5, 1.4, 3.8], font_size=8.0)
doc.add_heading("11.2 Cost control", level=2)
para("The current repository provides a component architecture but does not include a newly verified purchasing quote. The submission should present a dated BOM with vendor, part number, quantity, unit price, tax, shipping, lead time and substitute. Separate one channel proof cost, one sleeve alpha cost and two sleeve pilot cost. Do not reuse old prices without rechecking availability.")
add_table(["Cost bucket", "Contents", "Control"], [
    ("Electronics", "IMUs, haptic drivers, motors, hub, storage and power", "Freeze counts after architecture reconciliation"),
    ("Mechanical", "Printed housings, sleeve textile, straps and cable protection", "Prototype fit sizes before volume printing"),
    ("Tool interface", "Tabla pickups, loom switches or puppet sensor", "Validate only the flagship module first"),
    ("Testing", "Logic analyzer, fixtures, spare parts and participant materials", "Borrow laboratory instruments where appropriate"),
    ("Contingency", "Failed prints, replacement sensors and wiring", "Maintain explicit percentage and spares list"),
], widths=[1.35, 3.15, 2.25], font_size=8.0)

new_section("12 Sustainability accessibility and user experience")
doc.add_heading("12.1 Sustainability", level=2)
add_bullets([
    "Use replaceable modules and standard fasteners so failed electronics do not discard the complete sleeve.",
    "Prefer local fabrication for housings and textiles, with documented material selection and repair steps.",
    "Separate battery replacement and electronic waste handling from textile cleaning.",
    "Operate offline for practice and synchronize only approved data to reduce connectivity dependence.",
    "Measure useful device life and repair rate before claiming environmental benefit."
])
doc.add_heading("12.2 Accessibility and safety", level=2)
add_bullets([
    "Provide visual and audio equivalents for tactile cues where the learning goal allows.",
    "Support left handed and bilateral configurations without treating one orientation as default.",
    "Use clear cue intensity limits, skin contact inspection, emergency stop and battery protection.",
    "Avoid collecting disability or health information unless necessary, consented and protected.",
    "Include master and learner feedback in housing, cue mapping and terminology decisions."
])
doc.add_heading("12.3 User experience principles", level=2)
para("A successful interaction should require little setup, make calibration visible, explain why a cue occurred and show progress on cue free cycles. Learners should be able to pause, repeat, reduce intensity and delete a session. Masters should preview and approve the reference before learners can access it.")

new_section("13 Ethics data governance and cultural safeguards")
para("Embodied craft knowledge is not ordinary training data. Technical capability to record or classify a performance does not create permission to commercialize it. PARAMPARA should treat consent, attribution and benefit sharing as product requirements.")
add_table(["Control", "Minimum implementation"], [
    ("Consent", "Purpose, audience, duration, commercial use, withdrawal and future model use recorded before capture"),
    ("Provenance", "Master, session, phrase, instrument, sensor map, firmware and processing version retained"),
    ("Access", "Role based access with offline default and explicit export"),
    ("Approval", "Master reviews the reference package and learner facing representation"),
    ("Revocation", "Document how recordings, derived features and published references are removed"),
    ("Benefit", "Define attribution, compensation or community benefit before deployment"),
    ("Security", "Encrypt stored and transferred personal recordings; minimize retained raw data"),
], widths=[1.3, 5.3], font_size=8.1)
doc.add_heading("13.1 Prohibited claims and practices", level=2)
add_bullets([
    "Do not label a classifier score as authenticity, talent or cultural authority.",
    "Do not record masters or learners without informed, documented permission.",
    "Do not synthesize or imitate a named master's style outside the agreed use.",
    "Do not use generated photographs or simulations as evidence of a physical prototype.",
    "Do not imply endorsements, partnerships or field results that do not exist."
])

new_section("14 Impact adoption and scale")
doc.add_heading("14.1 Intended impact", level=2)
para("If validated, PARAMPARA could help institutions document movement based teaching, let learners practice between sessions and give masters more control over how their technique is represented. The near term impact claim is smaller: the project provides a testable engineering and research pathway for a problem that is often addressed only through media archives.")
doc.add_heading("14.2 Adoption pathway", level=2)
add_table(["Stage", "Partner type", "Offer", "Evidence needed"], [
    ("Lab pilot", "Institute electronics and design labs", "One channel and wearable alpha", "Bench timing, safety and repeatability"),
    ("Craft pilot", "Tabla teacher or cultural institution", "Consented capture and learner study", "Master approval and pilot protocol"),
    ("Institutional pilot", "Music school, museum or skilling centre", "Small cohort system and dashboard", "Usability, retention and support data"),
    ("Scale", "Training networks and cultural bodies", "Repairable kits and governed content library", "Cost, maintenance and multi site outcomes"),
], widths=[1.2, 1.75, 2.1, 1.55], font_size=7.8)
doc.add_heading("14.3 Business and operating model", level=2)
para("A credible early model is institution led rather than direct consumer hardware. Revenue may combine pilot kits, installation, maintenance, training and governed content services. Masters and partner communities should receive agreed attribution and benefit. Pricing, willingness to pay and procurement cycles remain unvalidated and require interviews and letters of intent.")

new_section("15 Risk register")
add_table(["Risk", "Likelihood", "Impact", "Mitigation and proof"], [
    ("Wearable alters technique", "Medium", "High", "Run E0 equivalence test; simplify site count if needed"),
    ("I2C bus or cable instability", "Medium", "High", "One channel then switched bus tests with error logs"),
    ("Haptic cues are late or confusing", "Medium", "High", "Measure physical onset and run blinded localization study"),
    ("Fingerprint learns instrument or session", "High", "High", "Grouped splits, refitting days and swapped instruments"),
    ("Learners depend on cues", "Medium", "High", "No cue checks and delayed device off test"),
    ("Battery safety or brownout", "Low to medium", "High", "Protected charging, load testing, thermal log and enclosure review"),
    ("Cultural extraction or misuse", "Medium", "High", "Master controlled consent, provenance, revocation and benefit sharing"),
    ("Scope becomes unbuildable", "High", "Medium", "Tabla first; gate Kathputli and handloom behind flagship validation"),
    ("Submission links fail", "Medium", "Medium", "Public signed out test, stable filenames and backup copy"),
], widths=[1.55, 0.85, 0.75, 3.55], font_size=7.5)

new_section("16 SIH evaluation alignment")
add_table(["Official criterion", "How PARAMPARA addresses it", "Evidence judges can inspect"], [
    ("Novelty", "Master governed body plus tool reference, confound controlled validation and device off assessment", "Architecture, protocol and prior art boundary"),
    ("Complexity", "Multi sensor routing, synchronization, haptics, embedded firmware and grouped evaluation", "CAD, firmware, circuit and technical roadmap"),
    ("Clarity and detail", "Claims are separated into implemented, tested, designed and proposed", "This dossier and claim evidence matrix"),
    ("Feasibility", "Risk reduced from one channel bench to wearable and field pilot", "CAD exports, Wokwi build, phased plan"),
    ("Practicability", "Open palm design, offline use and tool specific event channels", "Equipment viewer and use cases"),
    ("Sustainability", "Repairable modules, local fabrication and controlled data lifecycle", "Design rules and future lifecycle measures"),
    ("Scale of impact", "Institution led path from tabla to other movement intensive crafts", "Adoption roadmap with staged evidence"),
    ("User experience", "Visible calibration, controllable cues, no cue checks and master approval", "Phone prototype and interface requirements"),
    ("Future progression", "Defined H1 to H3 and E0 to E5 validation gates", "Exit criteria and risk register"),
], widths=[1.25, 3.15, 2.3], font_size=7.6)
doc.add_heading("16.1 Recommended jury demonstration order", level=2)
add_numbered([
    "State the craft knowledge gap and the narrow research question in 20 seconds.",
    "Rotate the full sleeve and reveal sensor, hub and haptic layers while clearly calling it CAD.",
    "Run the Wokwi control path and show one retained serial scenario.",
    "Run the phone practice loop through a cue free check.",
    "Show the controller benchmark and explain why it is software evidence rather than device accuracy.",
    "Close with the one channel physical milestone and the consent governed pilot."
])

new_section("17 Submission package and link checklist")
para("Use one public landing page as the primary jury link. The page should expose the final film, interactive 3D equipment, Wokwi circuit, research PDF and evidence index. Large archives may be stored separately, but every link must be tested while signed out.")
add_table(["Submission item", "Current location or action", "Final check"], [
    ("GitHub repository", "https://github.com/HarshDubey23/winner", "Public and main branch contains final files"),
    ("Vercel website", "TO BE FILLED AFTER DEPLOYMENT", "Root opens submission dashboard"),
    ("Wokwi circuit", "https://wokwi.com/projects/477036096886805505", "Runs without editor permissions"),
    ("Jury film", "submission/film-v3/PARAMPARA-4min-Jury-Film.mp4", "Plays on mobile and desktop with audio"),
    ("Technical video", "submission/video/PARAMPARA-Digital-Prototype.mp4", "Captions and labels are readable"),
    ("Research dossier", "submission/final/PARAMPARA_SIH_Submission_Research_and_Evidence_Dossier.pdf", "Correct team details and current PS title"),
    ("3D viewer", "showcase/index.html", "Models rotate and labels load"),
    ("Evidence index", "submission/docs/PARAMPARA-Proof-of-Work-Index.md", "No broken relative links"),
], widths=[1.35, 3.55, 1.8], font_size=7.6)
doc.add_heading("17.1 Mandatory before upload", level=2)
add_bullets([
    "Replace every TO BE FILLED field with verified team and institute details.",
    "Confirm PS 26214 title, organization, theme and hardware category against the live SIH portal.",
    "Use the current official SIH idea presentation template without changing prescribed headings or slide count.",
    "Check file size, format and deadline with the institute SPOC.",
    "Open every public link in a signed out browser and on a phone.",
    "Keep a local and cloud backup with a checksum and submission timestamp.",
    "Ensure every image, soundtrack and dataset has a source or licence record."
])

new_section("18 Claim and evidence matrix")
add_table(["Claim", "Evidence file", "Allowed wording", "Status"], [
    ("Equipment designed", "cad/out and showcase models", "We created inspectable digital CAD and housings", "Designed"),
    ("Teaching logic implemented", "firmware/src/teaching.h and tests", "The host tested controller implements bounded fading and check cycles", "Tested in software"),
    ("Virtual circuit runs", "Wokwi project and retained logs", "The one channel ESP32 substitute circuit compiles and executes", "Tested virtually"),
    ("Phone loop runs", "lite app and scenario outputs", "Scripted inputs produce expected practice scores and cue reduction", "Tested in browser"),
    ("Performer information exists in public data", "gmd ablation results", "Timing and velocity distinguish electronic drum performers in this experiment", "Method evidence"),
    ("Physical sleeve works", "No evidence yet", "The sleeve is proposed and ready for staged bench validation", "Unproven"),
    ("Learning improves", "Protocol only", "A randomized device off pilot is planned", "Unproven"),
    ("Three crafts validated", "Concept CAD only for extensions", "Tabla is first; Kathputli and handloom are planned", "Unproven"),
], widths=[1.45, 1.65, 2.85, 0.85], font_size=7.2)

new_section("19 References")
refs = [
    ("[1]", "Smart India Hackathon", "Guidelines for College SPOC SIH 2024", "https://sih.gov.in/letters/Guidelines-College-SPOC.pdf"),
    ("[2]", "Espressif Systems", "ESP32 S3 MINI 1 datasheet", "https://www.espressif.com/sites/default/files/documentation/esp32-s3-mini-1_mini-1u_datasheet_en.pdf"),
    ("[3]", "Bosch Sensortec", "BMI270 inertial measurement unit datasheet", "https://www.bosch-sensortec.com/media/boschsensortec/downloads/datasheets/bst-bmi270-ds000.pdf"),
    ("[4]", "Texas Instruments", "DRV2605L product documentation", "https://www.ti.com/product/DRV2605L"),
    ("[5]", "Texas Instruments", "DRV2605L ERM and LRA haptic driver evaluation guide", "https://www.ti.com/lit/ug/slou400/slou400.pdf"),
    ("[6]", "Microsoft Research", "Haptic feedback enhances force skill learning", "https://www.microsoft.com/en-us/research/publication/haptic-feedback-enhances-force-skill-learning/"),
    ("[7]", "Columbia University", "Haptic guidance benefits musical motor learning", "https://www.ee.columbia.edu/~grindlay/pubs/Haptics_2008.pdf"),
    ("[8]", "Google Magenta", "Groove MIDI Dataset", "https://magenta.withgoogle.com/datasets/groove"),
    ("[9]", "Wokwi", "Wokwi documentation", "https://docs.wokwi.com/"),
    ("[10]", "Wokwi", "MPU6050 simulator part reference", "https://docs.wokwi.com/parts/wokwi-mpu6050"),
    ("[11]", "GitHub Docs", "Creating a GitHub Pages site", "https://docs.github.com/en/pages/getting-started-with-github-pages/creating-a-github-pages-site"),
    ("[12]", "Global Indigenous Data Alliance", "CARE Principles for Indigenous Data Governance", "https://www.gida-global.org/care"),
    ("[13]", "Project repository", "PARAMPARA source code CAD simulations and evidence", "https://github.com/HarshDubey23/winner"),
    ("[14]", "Project simulation", "Public PARAMPARA Wokwi circuit", "https://wokwi.com/projects/477036096886805505"),
]
for tag, org, title, url in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.add_run(f"{tag} {org}. ").bold = True
    add_hyperlink(p, title, url)
    p.add_run(". Accessed 5 October 2026.")

doc.add_heading("Repository evidence consulted", level=2)
add_bullets([
    "PARAMPARA submission research evidence report and proof of work index.",
    "CAD sources, STEP and STL exports, rendered equipment views and browser models.",
    "Firmware teaching core, host tests and generated benchmark outputs.",
    "Wokwi sketch, circuit description, screenshots and retained serial logs.",
    "Phone prototype, scripted scenario outputs and export behavior.",
    "Groove MIDI experiment scripts, grouped evaluation outputs and ablation figure.",
    "Film, subtitles, media source registers and technical video metadata."
])

new_section("Appendix A Jury ready answer sheet")
add_table(["Likely question", "Evidence based answer"], [
    ("Where is the physical prototype?", "The current evidence is digital CAD, tested software and virtual electronics. The next milestone is a measured one sensor to motor bench channel before scaling to a wearable."),
    ("What is actually new?", "The combination of master governed capture, synchronized body and tool events, confound controlled fingerprint validation, bounded cue fading and device off assessment."),
    ("What does 47.6 percent mean?", "Balanced accuracy for identifying seven electronic drum performers across held sessions in a public MIDI experiment. It is not tabla or wearable accuracy."),
    ("Why not just use video?", "The project tests whether synchronized movement and tool events add repeatable information and useful cues. That benefit is a hypothesis to be tested against audio video practice."),
    ("Will learners become dependent?", "Regular no cue cycles and delayed device off tests are central evaluation requirements."),
    ("How do you protect masters?", "Consent, permitted use, provenance, approval, revocation and benefit sharing are designed as system requirements."),
    ("Why three crafts?", "Tabla is the validation case. Kathputli and handloom demonstrate transfer potential and remain gated until their sensing and outcomes are separately validated."),
    ("What will you build first?", "One timestamped sensor to ESP32 to LRA channel with measured latency, jitter, power and failure behavior."),
], widths=[2.0, 4.55], font_size=7.8)

new_section("Appendix B Final submission quality gate")
add_table(["Check", "Owner", "Status"], [
    ("Team and institute fields completed", "Team leader", "Pending"),
    ("Live SIH problem statement details verified", "SPOC and team leader", "Pending"),
    ("Official PPT template followed exactly", "Presentation owner", "Pending"),
    ("Research PDF opens and text is selectable", "Documentation owner", "Ready after export"),
    ("Vercel link loads while signed out", "Web owner", "Pending deployment"),
    ("Jury film plays with intelligible audio", "Media owner", "Ready locally"),
    ("3D viewer works on laptop and phone", "CAD and web owner", "Test before submission"),
    ("Wokwi project runs from public link", "Embedded owner", "Public link available"),
    ("Claims match evidence matrix", "All members", "Review required"),
    ("Source and licence records included", "Documentation owner", "Review required"),
    ("Backup package and checksum stored", "Team leader", "Pending final package"),
], widths=[3.65, 1.75, 1.1], font_size=8.0)

# Document properties and update fields
doc.core_properties.title = "PARAMPARA Research Engineering and Prototype Evidence Dossier"
doc.core_properties.subject = "Smart India Hackathon 2026 hardware submission supporting research and evidence"
doc.core_properties.author = "PARAMPARA Team"
doc.core_properties.keywords = "SIH, PARAMPARA, heritage, wearable, haptics, tabla, Kathputli, handloom"
settings = doc.settings._element
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)

doc.save(DOCX)
print(DOCX)
